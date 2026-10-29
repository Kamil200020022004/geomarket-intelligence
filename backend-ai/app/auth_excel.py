"""Temporary Excel authentication for local development only.
Not designed for concurrent servers or production deployment.
"""
from __future__ import annotations

import base64
import hashlib
import hmac
import json
import os
import secrets
import tempfile
import threading
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, Header, HTTPException, status
from pydantic import BaseModel, Field
from openpyxl import Workbook, load_workbook

router = APIRouter(prefix="/auth", tags=["Temporary Excel Auth"])
DATA_DIR = Path(__file__).resolve().parents[1] / "data"
USERS_PATH = DATA_DIR / "geomarket_users.xlsx"
SECRET_PATH = DATA_DIR / ".auth_secret"
LOCK = threading.RLock()
SESSION_SECONDS = 8 * 60 * 60


def get_secret() -> bytes:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not SECRET_PATH.exists():
        try:
            fd = os.open(SECRET_PATH, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            pass
        else:
            with os.fdopen(fd, "wb") as f:
                f.write(secrets.token_bytes(32))
    return SECRET_PATH.read_bytes()


def ensure_users_file():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not USERS_PATH.exists():
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Users"
        sheet.append(["user_id", "name", "email", "password_hash", "created_at"])
        wb.save(USERS_PATH)


def get_users():
    ensure_users_file()
    wb = load_workbook(USERS_PATH, read_only=True, data_only=True)
    try:
        sheet = wb["Users"]
        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] and row[2]:
                records.append(dict(zip(
                    ["user_id", "name", "email", "password_hash", "created_at"], row[:5]
                )))
        return records
    finally:
        wb.close()


def append_user(user):
    ensure_users_file()
    wb = load_workbook(USERS_PATH)
    try:
        wb["Users"].append([user[k] for k in [
            "user_id", "name", "email", "password_hash", "created_at"
        ]])
        with tempfile.NamedTemporaryFile(dir=DATA_DIR, suffix=".xlsx", delete=False) as f:
            tmp = Path(f.name)
        try:
            wb.save(tmp)
            os.replace(tmp, USERS_PATH)
        finally:
            tmp.unlink(missing_ok=True)
    finally:
        wb.close()


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    rounds = 600_000
    derived = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, rounds)
    return f"pbkdf2_sha256${rounds}${salt.hex()}${derived.hex()}"


def check_password(password: str, stored: str) -> bool:
    try:
        scheme, count, salt_hex, digest_hex = str(stored).split("$")
        if scheme != "pbkdf2_sha256":
            return False
        check = hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), int(count)
        )
        return hmac.compare_digest(check, bytes.fromhex(digest_hex))
    except (ValueError, TypeError):
        return False


def b64url(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b"=").decode("ascii")


def unb64url(value: str) -> bytes:
    return base64.urlsafe_b64decode(value + "=" * (-len(value) % 4))


def make_token(user_id: str) -> str:
    data = {"sub": user_id, "iat": int(time.time()), "exp": int(time.time()) + SESSION_SECONDS}
    payload = b64url(json.dumps(data, separators=(",", ":")).encode())
    signature = b64url(hmac.new(get_secret(), payload.encode(), hashlib.sha256).digest())
    return f"{payload}.{signature}"


def token_to_user_id(token: str) -> str:
    try:
        payload, signature = token.split(".")
        expected = b64url(hmac.new(get_secret(), payload.encode(), hashlib.sha256).digest())
        if not hmac.compare_digest(signature, expected):
            raise ValueError("invalid signature")
        data = json.loads(unb64url(payload))
        if data.get("exp", 0) <= int(time.time()) or not isinstance(data.get("sub"), str):
            raise ValueError("expired or invalid")
        return data["sub"]
    except (ValueError, TypeError, KeyError, json.JSONDecodeError):
        raise HTTPException(status_code=401, detail="Sesja wygasła. Zaloguj się ponownie.")


def public_user(user):
    return {"id": user["user_id"], "name": user["name"], "email": user["email"]}


class RegisterBody(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=8, max_length=128)


class LoginBody(BaseModel):
    email: str
    password: str


@router.post("/register", status_code=status.HTTP_201_CREATED)
def register(body: RegisterBody):
    email = body.email.strip().lower()
    if "@" not in email or "." not in email.rsplit("@", 1)[-1]:
        raise HTTPException(status_code=422, detail="Podaj poprawny adres e-mail.")
    with LOCK:
        if any(str(u["email"]).lower() == email for u in get_users()):
            raise HTTPException(status_code=409, detail="Konto z tym adresem już istnieje.")
        user = {
            "user_id": str(uuid.uuid4()),
            "name": body.name.strip(),
            "email": email,
            "password_hash": hash_password(body.password),
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        append_user(user)
    return {"user": public_user(user), "access_token": make_token(user["user_id"]), "token_type": "bearer"}


@router.post("/login")
def login(body: LoginBody):
    email = body.email.strip().lower()
    with LOCK:
        user = next((u for u in get_users() if str(u["email"]).lower() == email), None)
    if not user or not check_password(body.password, str(user["password_hash"])):
        raise HTTPException(status_code=401, detail="Nieprawidłowy e-mail lub hasło.")
    return {"user": public_user(user), "access_token": make_token(str(user["user_id"])), "token_type": "bearer"}


def current_user(authorization: str | None = Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Wymagane logowanie.")
    user_id = token_to_user_id(authorization[len("Bearer "):])
    with LOCK:
        user = next((u for u in get_users() if str(u["user_id"]) == user_id), None)
    if not user:
        raise HTTPException(status_code=401, detail="Nie znaleziono użytkownika.")
    return public_user(user)


@router.get("/me")
def me(user: dict = Depends(current_user)):
    return {"user": user}


@router.post("/logout")
def logout():
    # In this local prototype, logout removes token on the frontend.
    # Stateless tokens cannot be revoked before expiry without a denylist.
    return {"message": "Usuń token po stronie aplikacji."}
