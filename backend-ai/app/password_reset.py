
"""Password reset for GeoMarket Intelligence.

Password recovery using Resend and the existing Excel authentication system.
"""

from __future__ import annotations

import hashlib
import html
import json
import logging
import os
import secrets
import tempfile
import time

from pathlib import Path
from urllib.parse import quote

import httpx
from dotenv import load_dotenv

from fastapi import APIRouter, HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.auth_excel import (
    DATA_DIR,
    USERS_PATH,
    LOCK,
    get_users,
    hash_password,
)


# ============================================================
# AUTOMATYCZNE WCZYTYWANIE KONFIGURACJI
# ============================================================

# Lokalizacja pliku: backend-ai/.env
BACKEND_ROOT = Path(__file__).resolve().parents[1]

load_dotenv(BACKEND_ROOT / ".env")


# ============================================================
# KONFIGURACJA
# ============================================================

router = APIRouter(
    prefix="/auth",
    tags=["Password Reset"],
)

logger = logging.getLogger(__name__)

RESET_PATH = DATA_DIR / ".password_reset_tokens.json"

RESET_TTL_SECONDS = 15 * 60
RESET_COOLDOWN_SECONDS = 60

RESET_HISTORY: dict[str, float] = {}

PUBLIC_MESSAGE = (
    "Jeśli konto istnieje, wyślemy link do zmiany hasła. "
    "Sprawdź swoją pocztę."
)


# ============================================================
# MODELE DANYCH
# ============================================================

class ForgotPasswordBody(BaseModel):
    email: str = Field(
        min_length=3,
        max_length=254,
    )


class ResetPasswordBody(BaseModel):
    token: str = Field(
        min_length=20,
        max_length=300,
    )

    password: str = Field(
        min_length=8,
        max_length=128,
    )


# ============================================================
# ODCZYT TOKENÓW RESETOWANIA
# ============================================================

def _read_tokens() -> dict:
    if not RESET_PATH.exists():
        return {}

    try:
        data = json.loads(
            RESET_PATH.read_text(encoding="utf-8")
        )

        return data if isinstance(data, dict) else {}

    except (ValueError, OSError):
        logger.exception(
            "Nie udało się odczytać tokenów resetowania."
        )

        raise HTTPException(
            status_code=503,
            detail="Usługa resetowania jest chwilowo niedostępna.",
        )


# ============================================================
# ZAPIS TOKENÓW
# ============================================================

def _atomic_write_tokens(data: dict) -> None:
    DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with tempfile.NamedTemporaryFile(
        dir=DATA_DIR,
        suffix=".json",
        mode="w",
        encoding="utf-8",
        delete=False,
    ) as file:
        tmp = Path(file.name)
        json.dump(data, file)

    try:
        os.replace(tmp, RESET_PATH)

    finally:
        tmp.unlink(missing_ok=True)


# ============================================================
# WYSYŁANIE WIADOMOŚCI PRZEZ RESEND
# ============================================================

def _send_reset_email(email: str, link: str) -> None:
    api_key = os.environ.get(
        "RESEND_API_KEY", ""
    ).strip()

    sender = os.environ.get(
        "RESEND_FROM_EMAIL", ""
    ).strip()

    if not api_key or not sender:
        raise RuntimeError(
            "Brak konfiguracji RESEND_API_KEY lub RESEND_FROM_EMAIL."
        )

    safe_link = html.escape(
        link,
        quote=True,
    )

    payload = {
        "from": sender,
        "to": [email],
        "subject": "GeoMarket - resetowanie hasła",
        "html": (
            "<p>Otrzymaliśmy prośbę o zmianę hasła "
            "do konta GeoMarket Intelligence.</p>"
            f'<p><a href="{safe_link}">'
            "Ustaw nowe hasło"
            "</a></p>"
            "<p>Link jest ważny przez 15 minut "
            "i może zostać wykorzystany tylko raz.</p>"
            "<p>Jeżeli nie prosiłeś o zmianę hasła, "
            "zignoruj tę wiadomość.</p>"
        ),
        "text": (
            "Otrzymaliśmy prośbę o zmianę hasła.\n"
            f"Ustaw nowe hasło: {link}\n"
            "Link jest ważny przez 15 minut "
            "i może zostać wykorzystany tylko raz."
        ),
    }

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    logger.info(
        "Próba wysłania wiadomości resetującej."
    )

    try:
        with httpx.Client(
            timeout=15.0,
            trust_env=True,
            follow_redirects=False,
        ) as client:
            response = client.post(
                "https://api.resend.com/emails",
                headers=headers,
                json=payload,
            )

    except httpx.RequestError as error:
        logger.error(
            "Błąd połączenia z Resend: %s",
            type(error).__name__,
        )

        raise RuntimeError(
            "Nie udało się połączyć z Resend."
        ) from error

    if 200 <= response.status_code < 300:
        logger.warning(
            "SUKCES: Resend zaakceptował wiadomość. HTTP %s",
            response.status_code,
        )
        return

    logger.error(
        "RESEND ERROR: HTTP %s | %s",
        response.status_code,
        response.text[:2000],
    )

    raise RuntimeError(
        f"Resend rejected email: HTTP {response.status_code}"
    )


# ============================================================
# PROŚBA O ODZYSKANIE HASŁA
# POST /auth/forgot-password
# ============================================================

@router.post("/forgot-password")
def forgot_password(body: ForgotPasswordBody):
    email = body.email.strip().lower()

    with LOCK:
        now = time.time()

        if (
            now - RESET_HISTORY.get(email, 0)
            < RESET_COOLDOWN_SECONDS
        ):
            return {
                "message": PUBLIC_MESSAGE
            }

        RESET_HISTORY[email] = now

        user = next(
            (
                u
                for u in get_users()
                if str(u["email"]).strip().lower() == email
            ),
            None,
        )

        if user is None:
            return {
                "message": PUBLIC_MESSAGE
            }

        if (
            not os.environ.get("RESEND_API_KEY")
            or not os.environ.get("RESEND_FROM_EMAIL")
        ):
            logger.error(
                "Wysyłanie wiadomości resetującej "
                "nie zostało skonfigurowane."
            )

            return {
                "message": PUBLIC_MESSAGE
            }

        token = secrets.token_urlsafe(32)

        token_hash = hashlib.sha256(
            token.encode("utf-8")
        ).hexdigest()

        tokens = _read_tokens()

        tokens = {
            key: value
            for key, value in tokens.items()
            if (
                float(value.get("expires", 0)) > now
                and str(value.get("user_id"))
                != str(user["user_id"])
            )
        }

        tokens[token_hash] = {
            "user_id": str(user["user_id"]),
            "expires": now + RESET_TTL_SECONDS,
        }

        _atomic_write_tokens(tokens)

    # ========================================================
    # GENEROWANIE LINKU DO FORMULARZA
    # ========================================================

    frontend_url = os.environ.get(
        "FRONTEND_URL",
        "http://localhost:5173",
    ).rstrip("/")

    link = (
        f"{frontend_url}/reset-password"
        f"?token={quote(token)}"
    )

    # ========================================================
    # WYSYŁANIE WIADOMOŚCI
    # ========================================================

    try:
        _send_reset_email(email, link)

    except Exception:
        logger.exception(
            "Nie udało się wysłać wiadomości resetującej."
        )

        with LOCK:
            tokens = _read_tokens()
            tokens.pop(token_hash, None)
            _atomic_write_tokens(tokens)

    return {
        "message": PUBLIC_MESSAGE
    }


# ============================================================
# USTAWIENIE NOWEGO HASŁA
# POST /auth/reset-password
# ============================================================

@router.post("/reset-password")
def reset_password(body: ResetPasswordBody):
    token_hash = hashlib.sha256(
        body.token.encode("utf-8")
    ).hexdigest()

    with LOCK:
        tokens = _read_tokens()

        entry = tokens.get(token_hash)

        if (
            not entry
            or float(entry.get("expires", 0)) <= time.time()
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "Link jest nieprawidłowy lub wygasł. "
                    "Poproś o nowy."
                ),
            )

        user_id = str(entry["user_id"])

        users = get_users()

        if not any(
            str(user["user_id"]) == user_id
            for user in users
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "Link jest nieprawidłowy lub wygasł. "
                    "Poproś o nowy."
                ),
            )

        new_hash = hash_password(body.password)

        # ====================================================
        # ZAPIS NOWEGO HASŁA W EXCELU
        # ====================================================

        wb = load_workbook(USERS_PATH)

        try:
            sheet = wb["Users"]

            changed = False

            for row in sheet.iter_rows(min_row=2):
                if str(row[0].value) == user_id:
                    row[3].value = new_hash
                    changed = True
                    break

            if not changed:
                raise HTTPException(
                    status_code=400,
                    detail="Nie znaleziono konta.",
                )

            with tempfile.NamedTemporaryFile(
                dir=DATA_DIR,
                suffix=".xlsx",
                delete=False,
            ) as file:
                tmp = Path(file.name)

            try:
                wb.save(tmp)
                os.replace(tmp, USERS_PATH)

            finally:
                tmp.unlink(missing_ok=True)

        finally:
            wb.close()

        # Unieważnienie wykorzystanych tokenów resetowania.
        tokens = {
            key: value
            for key, value in tokens.items()
            if str(value.get("user_id")) != user_id
        }

        _atomic_write_tokens(tokens)

    return {
        "message": (
            "Hasło zostało zmienione. "
            "Możesz się zalogować."
        )
    }
